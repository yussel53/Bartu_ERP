# -*- coding: utf-8 -*-
"""dnm.xls vs Parametreler.xlsx vs Bartu paramDefs — net ayrım."""
from __future__ import annotations

import json
import re
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
import xlrd

BASE = Path(r"C:\Users\Yusuf\Desktop\EXCEL\Parametreler")
OUT = Path(r"d:\Dosyalar\YAZILIMLARIMIZ\Bartu_ERP\tools\_param_compare")
OUT.mkdir(parents=True, exist_ok=True)
DEFS = Path(r"d:\Dosyalar\YAZILIMLARIMIZ\Bartu_ERP\FrontendUI\src\data\paramDefs")


def norm(s: object) -> str:
    if s is None:
        return ""
    if isinstance(s, float) and s == int(s):
        s = str(int(s))
    else:
        s = str(s).strip()
    if s.startswith("+"):
        s = s[1:].strip()
    return re.sub(r"\s+", " ", s.replace("\u00a0", " ")).strip()


def fold(s: str) -> str:
    s = norm(s).casefold().replace("i\u0307", "i")
    s = s.translate(
        str.maketrans({"ı": "i", "ş": "s", "ğ": "g", "ü": "u", "ö": "o", "ç": "c"})
    )
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def read_xls(path: Path) -> list[list[str]]:
    book = xlrd.open_workbook(str(path))
    sh = book.sheet_by_index(0)
    return [[norm(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]


def read_xlsx_sheets(path: Path) -> dict[str, list[list[str]]]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = {}
    for s in wb.sheetnames:
        sh = wb[s]
        out[s] = [
            [norm(v) for v in row] for row in sh.iter_rows(values_only=True)
        ]
    return out


def is_module_header(name: str, value: str) -> bool:
    """Uzun tireli başlık: 'Malzeme Yönetimi Parametreleri ----'."""
    if value:
        return False
    n = norm(name)
    if n.count("-") >= 6 or "─" in n or "–" in n:
        return True
    if re.search(r"parametreleri\s*-{3,}", n, re.I):
        return True
    return False


def clean_header(name: str) -> str:
    return re.sub(r"[-–—─=\s]+$", "", norm(name)).strip()


def is_sub_header(name: str, value: str) -> bool:
    if value:
        return False
    f = fold(name)
    return f in (
        "ondegerler",
        "parametreler",
        "genel parametreler",
        "satinalma faturalari",
        "satis faturalari",
        "satinalma",
        "satis",
    )


def parse_dnm(rows: list[list[str]]) -> list[dict]:
    """dnm.xls: col0=ağaç işareti, col1=Adı, col2=Değeri."""
    items: list[dict] = []
    module = ""
    kind_block = ""  # Öndeğerler / Parametreler
    for r in rows:
        # + | Adı | Değeri  veya  Adı | Değeri
        if len(r) >= 3 and (r[0] in ("+", "-", "") or len(r[0]) <= 2):
            name, value = r[1], r[2]
        else:
            name, value = (r[0] if r else ""), (r[1] if len(r) > 1 else "")
        name = name.strip()
        value = value.strip()
        if not name or fold(name) in ("adi", "ad"):
            continue
        if is_module_header(name, value):
            module = clean_header(name)
            kind_block = ""
            continue
        if is_sub_header(name, value):
            kind_block = clean_header(name)
            continue
        items.append(
            {
                "module": module,
                "block": kind_block,
                "name": name,
                "value": value,
                "fold": fold(name),
            }
        )
    return items


def parse_excel_sheet(rows: list[list[str]], sheet: str) -> list[dict]:
    items: list[dict] = []
    section = sheet
    name_i, val_i = 0, 1
    for r in rows[:5]:
        folds = [fold(c) for c in r]
        if any(x in ("adi", "ad") for x in folds):
            for i, x in enumerate(folds):
                if x in ("adi", "ad"):
                    name_i = i
                if "deger" in x:
                    val_i = i
            break
    for r in rows:
        if len(r) <= max(name_i, val_i):
            continue
        name, value = r[name_i], r[val_i]
        if not name or fold(name) in ("adi", "ad"):
            continue
        if is_module_header(name, value) or is_sub_header(name, value):
            section = clean_header(name) or sheet
            continue
        if not value and len(name) < 60 and "parametre" in fold(name):
            section = clean_header(name)
            continue
        items.append(
            {
                "sheet": sheet,
                "section": section,
                "name": name,
                "value": value,
                "fold": fold(name),
            }
        )
    return items


def load_param_defs() -> list[dict]:
    found: list[dict] = []
    for p in DEFS.glob("*.ts"):
        if p.name in ("types.ts", "index.ts"):
            continue
        text = p.read_text(encoding="utf-8")
        for m in re.finditer(r"code:\s*'([^']+)'[\s\S]*?title:\s*'([^']+)'", text):
            code, title = m.group(1), m.group(2)
            found.append({"module": p.stem, "code": code, "title": title, "fold": fold(title)})
    return found


# Tam eşleşme / bilinen eş anlamlı (fold → code)
ALIASES: dict[str, str] = {
    "negatif stok seviyesi kontrolu": "STK_NEG_CTRL",
    "asgari stok seviyesi kontrolu": "STK_MIN_CTRL",
    "azami stok seviyesi kontrolu": "STK_MAX_CTRL",
    "stok maliyet turu": "STK_COST_METHOD",
    "malzeme fislerindeki bos muhasebe kodlarini otomatik doldur": "STK_FILL_ACC",
    "otomatik fis basimi": "STK_AUTO_PRINT",
    "fis satirlari ambar bilgisi degistirilebilsin": "STK_WHOUSE_EDIT",
    "stok seviye kontrolleri": "STK_LEVEL_BASE",
    "kullanilacak para birimi malzeme satir": "STK_LINE_CUR",
    "kullanilacak para birimi malzeme genel": "STK_HDR_CUR",
    "satis kdv orani": "STK_KDV_SAL",
    "alim kdv orani": "STK_KDV_PUR",
    "satinalma kdv orani": "STK_KDV_PUR",
    "iade kdv orani": "STK_KDV_RET",
    "satis iade kdv orani": "STK_KDV_RET",
    "alim iade kdv orani": "STK_KDV_RET",
    "faturalardaki bos muhasebe kodlarini otomatik doldur": "FAT_FILL_ACC",
    "faturalardaki bos smm muhasebe kodlarini otomatik doldur": "FAT_FILL_SMM",
    "odeme tutari fatura siparis tutarindan fazla olursa": "FAT_PAY_OVER",
    "fatura satirlarinda birim fiyat kontrolu": "FAT_PRICE_CTRL",
    "kullanilacak para birimi fatura genel": "FAT_CUR_HDR",
    "kullanilacak para birimi fatura satir": "FAT_CUR_LINE",
    "fatura durumu ondegeri": "FAT_STATUS_DEF",
    "tevkifat payi": "FAT_TEVK_PAY",
    "tevkifat paydasi": "FAT_TEVK_PAYDA",
    "indirimde yuvarlama": "FAT_DISC_ROUND",
    "faturalarda nakit odeme tahsilat fisi otomatik olarak acilsin": "FAT_OPEN_CASH",
    "fatura iptal edildiginde bagli irsaliye iptal edilsin": "FAT_CANCEL_DESP",
    "fatura satirlari": "FAT_LINES_PRINT",
    "otomatik fatura basimi": "FAT_AUTO_PRINT",
    "e faturalar basilsin": "FAT_EINV_PRINT",
    "birim fiyat ondegeri": "PUR_PRICE_DEF",  # bağlama satınalma/satış'a göre
    "birim fiyat kurus hanesi": "CAR_PRICE_DEC",
    "gunluk doviz kurus hanesi": "CAR_FX_DEC",
    "kdv hesaplamasinda yuvarlama sekli": "CAR_VAT_ROUND",
    "sipariste risk limiti asildiginda": "CAR_RISK_ORDER",
    "irsaliyede risk limiti asildiginda": "CAR_RISK_DESP",
    "acik hesaplarda risk limiti asildiginda": "CAR_RISK_OPEN",
    "cari hesap fislerindeki bos muhasebe kodlarini otomatik doldur": "CAR_FILL_ACC",
    "kullanilacak para birimi cari hesap satir": "CAR_LINE_CUR",
    "faturalarda cari bazinda belge no tekrarlandiginda": "CAR_DOCNO_INV",
}

SKIP_MODULES = (
    "bakım yönetimi",
    "bakim yonetimi",
    "iş akış",
    "is akis",
    "üretim",
    "uretim",
    "kalite",
    "kampanya",
    "e-mağaza",
    "e-magaza",
    "e-iş",
    "e-is",
    "talep yönetimi",
    "talep yonetimi",
    "teklif yönetimi",
    "teklif yonetimi",
    "sözleşme",
    "sozlesme",
    "dağıtım",
    "dagitim",
    "ithalat",
    "ihracat",
    "fason",
    "proje",
    "sabit kıymet",
    "sabit kiymet",
    "bordro",
    "personel",
    "bütçe",
    "butce",
)

SKIP_NAME = (
    "skonsol",
    "kets",
    "ökc",
    "okc",
    "e-mağaza",
    "e-magaza",
    "karma koli",
    "hızlı üretim",
    "hizli uretim",
    "abc kodu",
    "enflasyon",
    "fiyat endeksi",
    "kkk",
    "varyant kodu şablon",
)

# Öndeğer / oran — Bartu’ya eklenebilecek netler
ADD_NOW_PATTERNS = [
    (r"^sat[iı][sş]\s+kdv\s+oran", "STK_KDV_SAL", "stok", "Satış KDV Oranı (%)", "20"),
    (r"^al[iı]m\s+kdv\s+oran|^sat[iı]nalma\s+kdv\s+oran", "STK_KDV_PUR", "stok", "Satınalma KDV Oranı (%)", "20"),
    (r"iade\s+kdv\s+oran", "STK_KDV_RET", "stok", "İade KDV Oranı (%)", "20"),
    (r"^g[uü]venlik\s+stok\s+seviyesi\s+kontrol", "STK_SAFE_CTRL", "stok", "Güvenlik Stok Seviyesi Kontrolü", "continue"),
    (r"^giris\s+kontrol[uü]$|^giriş\s+kontrolü$", "STK_IN_CTRL", "stok", "Giriş Kontrolü", "continue"),
    (r"^[cç][iı]k[iı][sş]\s+kontrol", "STK_OUT_CTRL", "stok", "Çıkış Kontrolü", "continue"),
    (r"malzeme\s+fi[sş]leri\s+durumu\s*\(ambar.*[cç][iı]k", "STK_VOUCH_OUT_STATUS", "stok", "Malzeme Fişleri Durumu (Ambar Çıkış)", "0"),
    (r"malzeme\s+fi[sş]leri\s+durumu\s*\(ambar.*gir", "STK_VOUCH_IN_STATUS", "stok", "Malzeme Fişleri Durumu (Ambar Giriş)", "0"),
    (r"kullan[iı]lacak\s+para\s+birimi\s*-\s*malzeme\s*\(genel\)", "STK_HDR_CUR", "stok", "Kullanılacak Para Birimi - Malzeme (Genel)", "doc"),
    (r"fi[sş]\s+sat[iı]rlar[iı]\s+ambar\s+bilgisi\s+g[uü]ncellenirken\s+uyar", "STK_WHOUSE_WARN", "stok", "Fiş Satırları Ambar Bilgisi Güncellenirken Uyar", "1"),
    (r"tarih\s+de[gğ]i[sş]ikli[gğ]inde.*raporlama", "STK_FX_DATE_REP", "stok", "Tarih Değişikliğinde Raporlama Dövizi Kuru", "update"),
    (r"seri/?lot.*otomatik\s+pencere|otomatik\s+pencere\s+a[cç][iı]l[iı]m", "STK_SERIAL_POPUP", "stok", "Seri/Lot Miktar Değişince Otomatik Pencere", "1"),
]


def module_skipped(module: str) -> bool:
    f = fold(module)
    return any(fold(x) in f for x in SKIP_MODULES)


def name_skipped(name: str) -> bool:
    f = fold(name)
    return any(fold(x) in f for x in SKIP_NAME)


def unique_by_fold(items: list[dict]) -> list[dict]:
    od: OrderedDict[str, dict] = OrderedDict()
    for it in items:
        f = it["fold"]
        if f not in od:
            od[f] = {**it, "occurrences": 1, "modules": [it.get("module") or it.get("sheet") or ""]}
        else:
            od[f]["occurrences"] += 1
            m = it.get("module") or it.get("sheet") or ""
            if m and m not in od[f]["modules"]:
                od[f]["modules"].append(m)
    return list(od.values())


def match_add_now(name: str) -> dict | None:
    for pat, code, mod, title, default in ADD_NOW_PATTERNS:
        if re.search(pat, name, re.I):
            return {"code": code, "module": mod, "title": title, "defaultValue": default, "from_name": name}
    return None


def main() -> None:
    dnm_rows = read_xls(BASE / "dnm.xls")
    dnm_items = parse_dnm(dnm_rows)
    dnm_unique = unique_by_fold(dnm_items)

    px = read_xlsx_sheets(BASE / "Parametreler.xlsx")
    excel_items: list[dict] = []
    for sheet, rows in px.items():
        excel_items.extend(parse_excel_sheet(rows, sheet))
    excel_unique = unique_by_fold(excel_items)
    excel_folds = {x["fold"] for x in excel_unique}

    defs = load_param_defs()
    defs_by_fold = {d["fold"]: d for d in defs}
    defs_by_code = {d["code"]: d for d in defs}

    buckets: dict[str, list[dict]] = defaultdict(list)
    add_now: list[dict] = []
    add_now_codes: set[str] = set()

    for it in dnm_unique:
        f = it["fold"]
        name = it["name"]
        code = ALIASES.get(f) or (defs_by_fold[f]["code"] if f in defs_by_fold else None)

        if code and code in defs_by_code:
            buckets["ayni_bartu"].append({**it, "mapped_code": code})
            continue

        if f in excel_folds:
            buckets["ayni_excel"].append(it)
            # öndeğer adayı mı — ekle?
            m = match_add_now(name)
            if m and m["code"] not in defs_by_code and m["code"] not in add_now_codes:
                add_now.append({**m, "dnm_value": it["value"], "also_in_excel": True})
                add_now_codes.add(m["code"])
            continue

        if module_skipped(it.get("module") or "") or name_skipped(name):
            buckets["kapsam_disi"].append(it)
            continue

        m = match_add_now(name)
        if m:
            if m["code"] in defs_by_code:
                buckets["ayni_bartu"].append({**it, "mapped_code": m["code"]})
            else:
                if m["code"] not in add_now_codes:
                    add_now.append({**m, "dnm_value": it["value"], "also_in_excel": False})
                    add_now_codes.add(m["code"])
                buckets["ekle_simdi"].append({**it, "proposed": m})
            continue

        # Öndeğer bloğundaki sayısal/oran/durum — sorulacak öndeğer
        block_f = fold(it.get("block") or "")
        if block_f == "ondegerler" or any(
            x in f
            for x in (
                "kdv oran",
                "ondeger",
                "para birimi",
                "durumu",
                "kurus hanesi",
                "tevkifat",
                "stopaj",
            )
        ):
            buckets["sor_ondeger"].append(it)
            continue

        buckets["sor_param"].append(it)

    # Modül sayıları
    mod_counts = defaultdict(int)
    for it in dnm_unique:
        mod_counts[it.get("module") or "(modül yok)"] += 1

    report = {
        "summary": {
            "dnm_rows": len(dnm_items),
            "dnm_unique": len(dnm_unique),
            "excel_unique": len(excel_unique),
            "bartu_defs": len(defs),
            "ayni_bartu": len(buckets["ayni_bartu"]),
            "ayni_excel": len(buckets["ayni_excel"]),
            "ekle_simdi": len(add_now),
            "sor_ondeger": len(buckets["sor_ondeger"]),
            "sor_param": len(buckets["sor_param"]),
            "kapsam_disi": len(buckets["kapsam_disi"]),
        },
        "modules": dict(sorted(mod_counts.items(), key=lambda x: -x[1])),
        "add_now": add_now,
        "buckets": {k: v for k, v in buckets.items()},
        "bartu_defs": defs,
    }
    (OUT / "dnm_compare.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    lines = [
        "# dnm.xls — ayrım özeti",
        "",
        f"Kaynak: `dnm.xls` (**{report['summary']['dnm_unique']}** benzersiz) · "
        f"`Parametreler.xlsx` (**{report['summary']['excel_unique']}**) · "
        f"Bartu defs (**{report['summary']['bartu_defs']}**)",
        "",
        "## Sayılar",
        f"- Aynı (Bartu’da var): **{report['summary']['ayni_bartu']}**",
        f"- Aynı (Excel’de var, Bartu seed’inde yok / zaten planlı): **{report['summary']['ayni_excel']}**",
        f"- Şimdi eklenecek (öndeğer/kontrol): **{report['summary']['ekle_simdi']}**",
        f"- Sorulacak — öndeğer: **{report['summary']['sor_ondeger']}**",
        f"- Sorulacak — parametre: **{report['summary']['sor_param']}**",
        f"- Kapsam dışı: **{report['summary']['kapsam_disi']}**",
        "",
        "## dnm modül başlıkları",
        "",
    ]
    for m, c in report["modules"].items():
        lines.append(f"- {m}: {c}")

    lines += ["", "## Şimdi eklenen / eklenecek", ""]
    if not add_now:
        lines.append("_yok_")
    for a in add_now:
        lines.append(
            f"- `{a['code']}` · **{a['title']}** · def=`{a['defaultValue']}` · dnm=`{a.get('dnm_value','')}` · mod={a['module']}"
        )

    def dump(title: str, key: str, limit: int = 80) -> None:
        lines.append("")
        lines.append(f"## {title}")
        lines.append("")
        items = buckets.get(key) or []
        if not items:
            lines.append("_yok_")
            return
        for it in items[:limit]:
            mods = ", ".join(it.get("modules") or [it.get("module") or ""])
            code = f" → `{it['mapped_code']}`" if it.get("mapped_code") else ""
            val = f" = `{it['value']}`" if it.get("value") else ""
            lines.append(f"- **{it['name']}**{val}{code} · _{mods}_")
        if len(items) > limit:
            lines.append(f"- … +{len(items)-limit} (JSON)")

    dump("1) Aynı — Bartu", "ayni_bartu", 100)
    dump("2) Aynı — Parametreler.xlsx (Bartu’da kod yok / planlı)", "ayni_excel", 100)
    dump("3) Sorulacak — öndeğer", "sor_ondeger", 120)
    dump("4) Sorulacak — parametre", "sor_param", 120)
    dump("5) Kapsam dışı", "kapsam_disi", 40)

    (OUT / "dnm_compare.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    print("add_now", len(add_now))
    for a in add_now:
        print(" ", a["code"], a["title"], a.get("dnm_value"))
    print("wrote", OUT / "dnm_compare.md")


if __name__ == "__main__":
    main()
